import os
import re
import hashlib
from pathlib import Path

import streamlit as st
from pptx import Presentation
import pythoncom
import win32com.client
from google import genai
from PIL import Image
import fitz  # PyMuPDF

# =========================
# INSTELLINGEN
# =========================

GEMINI_API_KEY = "AIzaSyAyPQD388Qg977jKQ3X5ZSD8R_-09p86UM"
MODEL_NAME = "gemini-2.5-flash"

BASE_CACHE_DIR = Path("cache")
UPLOAD_CACHE_DIR = BASE_CACHE_DIR / "uploads"
IMG_CACHE_DIR = BASE_CACHE_DIR / "page_images"

UPLOAD_CACHE_DIR.mkdir(parents=True, exist_ok=True)
IMG_CACHE_DIR.mkdir(parents=True, exist_ok=True)

st.set_page_config(
    page_title="StudyCopilot",
    layout="wide",
    initial_sidebar_state="collapsed"
)

client = None
if GEMINI_API_KEY and GEMINI_API_KEY != "JOUW_GEMINI_KEY_HIER":
    client = genai.Client(api_key=GEMINI_API_KEY)

# =========================
# SESSION STATE
# =========================

defaults = {
    "theme_mode": "Light",
    "language": "Nederlands",
    "page_index": 0,
    "ai_outputs": {},
    "chat_answers": {},
    "show_page_overview": False,
    "current_file_marker": None,
    "uploaded_name": None,
    "uploaded_bytes": None,
}
for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v


# =========================
# THEME / CSS
# =========================

def get_theme_tokens(theme_mode: str):
    if theme_mode == "Dark":
        return {
            "bg": "#0b1220",
            "bg2": "#111827",
            "card": "#111827",
            "card2": "#1f2937",
            "text": "#f8fafc",
            "muted": "#cbd5e1",
            "border": "rgba(148, 163, 184, 0.18)",
            "shadow": "0 14px 34px rgba(0, 0, 0, 0.35)",
            "accent": "#60a5fa",
            "accent_soft": "rgba(96, 165, 250, 0.18)",
            "button_bg": "#1e293b",
            "button_bg_hover": "#334155",
            "button_text": "#f8fafc",
            "input_bg": "#0f172a",
            "code_bg": "#0f172a",
            "code_text": "#e5e7eb",
            "sidebar_bg": "#0f172a",
        }
    return {
        "bg": "#eaf4ff",
        "bg2": "#ffffff",
        "card": "#ffffff",
        "card2": "#ffffff",
        "text": "#0f172a",
        "muted": "#475569",
        "border": "rgba(148, 163, 184, 0.20)",
        "shadow": "0 10px 30px rgba(15, 23, 42, 0.08)",
        "accent": "#3b82f6",
        "accent_soft": "rgba(59, 130, 246, 0.10)",
        "button_bg": "#eff6ff",
        "button_bg_hover": "#dbeafe",
        "button_text": "#0f172a",
        "input_bg": "#ffffff",
        "code_bg": "#f8fafc",
        "code_text": "#0f172a",
        "sidebar_bg": "#ffffff",
    }


def apply_custom_css(theme_mode: str):
    t = get_theme_tokens(theme_mode)

    st.markdown(
        f"""
        <style>
            header[data-testid="stHeader"] {{
                background: transparent;
            }}

            .stApp {{
                background: linear-gradient(180deg, {t["bg"]} 0%, {t["bg2"]} 100%);
                color: {t["text"]};
            }}

            [data-testid="stSidebar"] {{
                background: {t["sidebar_bg"]} !important;
                border-right: 1px solid {t["border"]};
            }}

            .block-container {{
                max-width: 1100px;
                padding-top: 1rem;
                padding-bottom: 2rem;
            }}

            .hero-wrap {{
                min-height: 78vh;
                display: flex;
                align-items: center;
                justify-content: center;
            }}

            .hero-card {{
                width: 100%;
                max-width: 760px;
                margin: 0 auto;
                background: {t["card"]};
                border: 1px solid {t["border"]};
                border-radius: 24px;
                padding: 2rem;
                box-shadow: {t["shadow"]};
            }}

            .app-title {{
                font-size: 2.2rem;
                font-weight: 800;
                color: {t["text"]};
                text-align: center;
                margin-bottom: 0.4rem;
            }}

            .app-subtitle {{
                color: {t["muted"]};
                text-align: center;
                font-size: 1rem;
                margin-bottom: 1.5rem;
            }}

            .section-card {{
                background: {t["card"]};
                border: 1px solid {t["border"]};
                border-radius: 22px;
                padding: 1rem;
                box-shadow: {t["shadow"]};
            }}

            .accent-badge {{
                display: inline-block;
                background: {t["accent_soft"]};
                color: {t["accent"]};
                padding: 0.35rem 0.7rem;
                border-radius: 999px;
                font-weight: 700;
                font-size: 0.92rem;
                margin-bottom: 0.8rem;
            }}

            .small-note {{
                color: {t["muted"]};
                font-size: 0.92rem;
            }}

            /* File uploader */
            div[data-testid="stFileUploader"] {{
                background: {t["card"]} !important;
                border: 1px solid {t["border"]} !important;
                border-radius: 16px !important;
                padding: 0.65rem !important;
            }}

            /* Inputs */
            div[data-testid="stSelectbox"] > div,
            div[data-testid="stNumberInput"] input,
            div[data-testid="stTextInput"] input,
            textarea {{
                background: {t["input_bg"]} !important;
                color: {t["text"]} !important;
                border-radius: 12px !important;
            }}

            label, p, span, div {{
                color: {t["text"]};
            }}

            /* Buttons */
            div[data-testid="stButton"] > button,
            div[data-testid="stFormSubmitButton"] > button {{
                background: {t["button_bg"]} !important;
                color: {t["button_text"]} !important;
                border: 1px solid {t["border"]} !important;
                border-radius: 12px !important;
                font-weight: 700 !important;
                min-height: 44px !important;
                box-shadow: none !important;
            }}

            div[data-testid="stButton"] > button:hover,
            div[data-testid="stFormSubmitButton"] > button:hover {{
                background: {t["button_bg_hover"]} !important;
                color: {t["button_text"]} !important;
            }}

            div[data-testid="stButton"] > button:disabled,
            div[data-testid="stFormSubmitButton"] > button:disabled {{
                opacity: 0.55 !important;
                color: {t["button_text"]} !important;
            }}

            /* Markdown / code / formula blocks */
            pre, code, .stCodeBlock, .stCode {{
                background: {t["code_bg"]} !important;
                color: {t["code_text"]} !important;
                border-radius: 12px !important;
                border: 1px solid {t["border"]} !important;
            }}

            .stMarkdown pre, .stMarkdown code {{
                background: {t["code_bg"]} !important;
                color: {t["code_text"]} !important;
            }}

            .katex, .katex-display, .math, mjx-container {{
                background: transparent !important;
                color: {t["text"]} !important;
            }}

            [data-testid="stImage"] img {{
                border-radius: 18px;
            }}

            /* Mobile spacing */
            @media (max-width: 768px) {{
                .hero-card {{
                    padding: 1.2rem;
                    border-radius: 18px;
                }}

                .app-title {{
                    font-size: 1.7rem;
                }}

                .block-container {{
                    padding-top: 0.6rem;
                    padding-bottom: 1.4rem;
                }}
            }}
        </style>
        """,
        unsafe_allow_html=True
    )


apply_custom_css(st.session_state.theme_mode)

# =========================
# HULPFUNCTIES
# =========================

LANGUAGE_PROMPTS = {
    "Nederlands": "Antwoord volledig in het Nederlands.",
    "English": "Answer fully in English.",
    "Deutsch": "Antworte vollständig auf Deutsch.",
    "Français": "Réponds entièrement en français.",
    "Español": "Responde completamente en español."
}


def current_language_instruction():
    return LANGUAGE_PROMPTS.get(st.session_state.language, LANGUAGE_PROMPTS["Nederlands"])


def natural_sort_key(filename: str):
    return [int(text) if text.isdigit() else text.lower() for text in re.split(r"(\d+)", filename)]


def file_hash(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


def save_uploaded_once(file_bytes: bytes, hash_value: str, suffix: str) -> Path:
    file_path = UPLOAD_CACHE_DIR / f"{hash_value}{suffix}"
    if not file_path.exists():
        file_path.write_bytes(file_bytes)
    return file_path.resolve()


def export_pptx_to_images_if_needed(pptx_file_path: Path, hash_value: str) -> Path:
    output_folder = (IMG_CACHE_DIR / hash_value).resolve()
    output_folder.mkdir(parents=True, exist_ok=True)

    existing_pngs = sorted(
        [f for f in os.listdir(output_folder) if f.lower().endswith(".png")],
        key=natural_sort_key
    )
    if existing_pngs:
        return output_folder

    pythoncom.CoInitialize()
    powerpoint = None
    presentation = None

    try:
        powerpoint = win32com.client.Dispatch("PowerPoint.Application")
        powerpoint.Visible = 1

        pptx_file_path_str = str(pptx_file_path.resolve())
        output_folder_str = str(output_folder.resolve())

        presentation = powerpoint.Presentations.Open(pptx_file_path_str, False, False, False)
        presentation.SaveAs(output_folder_str, 18)  # PNG
        presentation.Close()
        presentation = None
        powerpoint.Quit()
        powerpoint = None

    except Exception as e:
        if presentation is not None:
            try:
                presentation.Close()
            except Exception:
                pass
        if powerpoint is not None:
            try:
                powerpoint.Quit()
            except Exception:
                pass

        raise RuntimeError(
            f"PowerPoint kon het bestand niet openen.\n\n"
            f"Bestandspad: {pptx_file_path}\n"
            f"Outputmap: {output_folder}\n"
            f"Fout: {e}"
        )

    return output_folder


def render_pdf_to_images_if_needed(pdf_path: Path, hash_value: str) -> Path:
    output_folder = (IMG_CACHE_DIR / hash_value).resolve()
    output_folder.mkdir(parents=True, exist_ok=True)

    existing_pngs = sorted(
        [f for f in os.listdir(output_folder) if f.lower().endswith(".png")],
        key=natural_sort_key
    )
    if existing_pngs:
        return output_folder

    doc = fitz.open(pdf_path)
    for i in range(len(doc)):
        page = doc.load_page(i)
        pix = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
        out = output_folder / f"page_{i + 1}.png"
        pix.save(str(out))
    doc.close()
    return output_folder


@st.cache_data(show_spinner=False)
def load_pptx_texts(pptx_path_str: str):
    presentation = Presentation(pptx_path_str)
    texts = []
    for slide in presentation.slides:
        slide_text = ""
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                slide_text += shape.text + "\n"
        texts.append(slide_text.strip())
    return texts


@st.cache_data(show_spinner=False)
def load_pdf_texts(pdf_path_str: str):
    doc = fitz.open(pdf_path_str)
    texts = []
    for i in range(len(doc)):
        page = doc.load_page(i)
        texts.append(page.get_text("text").strip())
    doc.close()
    return texts


def get_image_files(output_folder: Path):
    return sorted(
        [f for f in os.listdir(output_folder) if f.lower().endswith(".png")],
        key=natural_sort_key
    )


def build_deck_memory(page_texts, max_chars=18000):
    parts = []
    total_len = 0

    for i, txt in enumerate(page_texts):
        clean = txt.strip()
        if not clean:
            continue

        snippet = clean[:700]
        block = f"[Page {i + 1}]\n{snippet}\n\n"

        if total_len + len(block) > max_chars:
            remaining = max_chars - total_len
            if remaining > 150:
                parts.append(block[:remaining])
            break

        parts.append(block)
        total_len += len(block)

    return "".join(parts).strip() if parts else "No readable text in this file."


def get_cache_key(hash_value: str, page_index: int, kind: str, language: str):
    return f"{hash_value}__page_{page_index}__{kind}__{language}"


def pil_image_from_path(image_path: Path):
    return Image.open(image_path)


def call_gemini_multimodal(prompt: str, image_path: Path) -> str:
    if client is None:
        return "Gemini API key ontbreekt. Zet eerst je GEMINI_API_KEY goed."

    try:
        img = pil_image_from_path(image_path)
        response = client.models.generate_content(
            model=MODEL_NAME,
            contents=[img, prompt]
        )
        if hasattr(response, "text") and response.text:
            return response.text.strip()
        return "Geen antwoord ontvangen van Gemini."
    except Exception as e:
        return f"Gemini is tijdelijk niet beschikbaar. Probeer opnieuw.\n\nTechnische melding: {e}"


def generate_short_explanation(current_text: str, deck_memory: str, page_number: int, image_path: Path) -> str:
    prompt = f"""
{current_language_instruction()}

You are looking at the current page/slide image. Use the image as the main source.
Use extracted text and full-file memory only as support.

IMPORTANT:
- explain this page briefly
- maximum 5 sentences in the explanation
- then maximum 3 bullet points with what matters most
- do NOT keep referring to earlier pages/slides
- only mention another page if truly necessary
- if there is a formula, process, schema, arrows, graph, or structure, explain what it means
- make it useful for studying
- keep it compact and clear

Return exactly in this structure:

Short explanation:
...

Important:
- ...
- ...
- ...

Page number:
{page_number}

Current extracted text:
{current_text}

Memory of the full file:
{deck_memory}
"""
    return call_gemini_multimodal(prompt, image_path)


def generate_simpler_explanation(current_text: str, deck_memory: str, page_number: int, image_path: Path) -> str:
    prompt = f"""
{current_language_instruction()}

Use the current page image as the main source.

Explain this page in a much simpler way.

IMPORTANT:
- maximum 4 sentences
- very easy language
- short and direct
- do not keep mentioning earlier pages/slides

Page number:
{page_number}

Current extracted text:
{current_text}

Memory of the full file:
{deck_memory}
"""
    return call_gemini_multimodal(prompt, image_path)


def generate_key_points(current_text: str, deck_memory: str, page_number: int, image_path: Path) -> str:
    prompt = f"""
{current_language_instruction()}

Use the current page image as the main source.

Give only the most important takeaways from this page.

IMPORTANT:
- maximum 5 bullet points
- each bullet short
- no long explanation
- avoid mentioning other pages unless really necessary

Page number:
{page_number}

Current extracted text:
{current_text}

Memory of the full file:
{deck_memory}
"""
    return call_gemini_multimodal(prompt, image_path)


def generate_questions(current_text: str, deck_memory: str, page_number: int, image_path: Path) -> str:
    prompt = f"""
{current_language_instruction()}

Use the current page image as the main source.

Create 3 short study questions about this page.

IMPORTANT:
- exactly 3 questions
- put a short answer directly under each question
- compact
- do not keep referring to other pages unless really necessary

Page number:
{page_number}

Current extracted text:
{current_text}

Memory of the full file:
{deck_memory}
"""
    return call_gemini_multimodal(prompt, image_path)


def ask_about_page(current_text: str, deck_memory: str, page_number: int, image_path: Path, question: str, recent_chat_history: list):
    history_text = ""
    if recent_chat_history:
        for item in recent_chat_history[-6:]:
            history_text += f"User: {item['question']}\nAssistant: {item['answer']}\n\n"

    prompt = f"""
{current_language_instruction()}

You are an AI study coach.
Use the current page image as the main source.
Use the extracted text and file memory only as support.

IMPORTANT:
- answer only the user's question
- keep it concise
- maximum 7 sentences
- do not keep referring to other pages unless truly necessary
- if the question is about a formula, arrow, process, graph, or layout, explain it concretely

Recent chat:
{history_text}

Page number:
{page_number}

Current extracted text:
{current_text}

Memory of the full file:
{deck_memory}

User question:
{question}
"""
    return call_gemini_multimodal(prompt, image_path)


def process_uploaded_file(file_name: str, file_bytes: bytes):
    hash_value = file_hash(file_bytes)
    suffix = Path(file_name).suffix.lower()

    if suffix == ".pptx":
        file_path = save_uploaded_once(file_bytes, hash_value, ".pptx")
        output_folder = export_pptx_to_images_if_needed(file_path, hash_value)
        page_texts = load_pptx_texts(str(file_path))
        file_type_label = "dia"
    elif suffix == ".pdf":
        file_path = save_uploaded_once(file_bytes, hash_value, ".pdf")
        output_folder = render_pdf_to_images_if_needed(file_path, hash_value)
        page_texts = load_pdf_texts(str(file_path))
        file_type_label = "pagina"
    else:
        raise ValueError("Alleen .pptx en .pdf worden ondersteund.")

    image_files = get_image_files(output_folder)

    return {
        "hash_value": hash_value,
        "file_path": file_path,
        "output_folder": output_folder,
        "page_texts": page_texts,
        "image_files": image_files,
        "file_type_label": file_type_label,
    }


def store_uploaded_file(uploaded_file):
    if uploaded_file is None:
        return

    st.session_state.uploaded_name = uploaded_file.name
    st.session_state.uploaded_bytes = uploaded_file.getvalue()
    st.session_state.page_index = 0
    st.session_state.ai_outputs = {}
    st.session_state.chat_answers = {}
    st.session_state.show_page_overview = False
    st.session_state.current_file_marker = None


# =========================
# STARTSCHERM
# =========================

has_uploaded_file = (
    st.session_state.uploaded_name is not None and
    st.session_state.uploaded_bytes is not None
)

if not has_uploaded_file:
    st.markdown('<div class="hero-wrap">', unsafe_allow_html=True)
    st.markdown('<div class="hero-card">', unsafe_allow_html=True)
    st.markdown('<div class="app-title">StudyCopilot</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="app-subtitle">Upload je PowerPoint of PDF en krijg compacte AI-uitleg per dia of pagina.</div>',
        unsafe_allow_html=True
    )

    uploaded_file_home = st.file_uploader(
        "Upload PowerPoint of PDF",
        type=["pptx", "pdf"],
        key="home_uploader"
    )

    col_a, col_b = st.columns(2)

    with col_a:
        new_language = st.selectbox(
            "Taal",
            ["Nederlands", "English", "Deutsch", "Français", "Español"],
            index=["Nederlands", "English", "Deutsch", "Français", "Español"].index(st.session_state.language),
            key="home_language"
        )
        if new_language != st.session_state.language:
            st.session_state.language = new_language
            st.rerun()

    with col_b:
        new_theme = st.selectbox(
            "Thema",
            ["Light", "Dark"],
            index=["Light", "Dark"].index(st.session_state.theme_mode),
            key="home_theme"
        )
        if new_theme != st.session_state.theme_mode:
            st.session_state.theme_mode = new_theme
            st.rerun()

    if uploaded_file_home is not None:
        store_uploaded_file(uploaded_file_home)
        st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)
    st.markdown('</div>', unsafe_allow_html=True)
    st.stop()


# =========================
# SIDEBAR NA UPLOAD
# =========================

with st.sidebar:
    st.markdown("## Instellingen")

    uploaded_file_sidebar = st.file_uploader(
        "Upload nieuw bestand",
        type=["pptx", "pdf"],
        key="sidebar_uploader"
    )
    if uploaded_file_sidebar is not None:
        store_uploaded_file(uploaded_file_sidebar)
        st.rerun()

    new_language = st.selectbox(
        "Taal",
        ["Nederlands", "English", "Deutsch", "Français", "Español"],
        index=["Nederlands", "English", "Deutsch", "Français", "Español"].index(st.session_state.language),
        key="sidebar_language"
    )
    if new_language != st.session_state.language:
        st.session_state.language = new_language
        st.rerun()

    new_theme = st.selectbox(
        "Thema",
        ["Light", "Dark"],
        index=["Light", "Dark"].index(st.session_state.theme_mode),
        key="sidebar_theme"
    )
    if new_theme != st.session_state.theme_mode:
        st.session_state.theme_mode = new_theme
        st.rerun()

# =========================
# APP
# =========================

st.markdown('<div class="app-title" style="text-align:left; font-size:1.7rem; margin-bottom:0.2rem;">StudyCopilot</div>', unsafe_allow_html=True)
st.markdown('<div class="app-subtitle" style="text-align:left;">Begrijp je dia’s en pagina’s sneller met AI.</div>', unsafe_allow_html=True)

processed = process_uploaded_file(
    file_name=st.session_state.uploaded_name,
    file_bytes=st.session_state.uploaded_bytes
)

hash_value = processed["hash_value"]
output_folder = processed["output_folder"]
page_texts = processed["page_texts"]
image_files = processed["image_files"]
file_type_label = processed["file_type_label"]

current_file_marker = f"{st.session_state.uploaded_name}__{hash_value}"

if st.session_state.current_file_marker != current_file_marker:
    st.session_state.current_file_marker = current_file_marker
    st.session_state.page_index = 0
    st.session_state.ai_outputs = {}
    st.session_state.chat_answers = {}
    st.session_state.show_page_overview = False

total_pages = len(page_texts)
current_index = st.session_state.page_index
current_text = page_texts[current_index] if current_index < len(page_texts) else ""
deck_memory = build_deck_memory(page_texts, max_chars=18000)

image_path = None
if current_index < len(image_files):
    image_path = output_folder / image_files[current_index]

st.markdown(
    f'<div class="accent-badge">{file_type_label.capitalize()} {current_index + 1} van {total_pages}</div>',
    unsafe_allow_html=True
)

# Overzicht knop
top_left, top_right = st.columns([1, 1])

with top_left:
    if st.button(f"Overzicht {file_type_label}'s", use_container_width=True):
        st.session_state.show_page_overview = not st.session_state.show_page_overview
        st.rerun()

with top_right:
    st.empty()

# Overzicht grid
if st.session_state.show_page_overview:
    st.markdown('<div class="section-card">', unsafe_allow_html=True)
    st.markdown(f"### Overzicht van alle {file_type_label}'s")

    cols_per_row = 4
    for row_start in range(0, len(image_files), cols_per_row):
        cols = st.columns(cols_per_row)
        batch = image_files[row_start:row_start + cols_per_row]

        for idx, image_file in enumerate(batch):
            page_number = row_start + idx
            thumb_path = output_folder / image_file

            with cols[idx]:
                st.image(str(thumb_path), use_container_width=True)
                if st.button(
                    f"{file_type_label.capitalize()} {page_number + 1}",
                    key=f"thumb_{page_number}",
                    use_container_width=True
                ):
                    st.session_state.page_index = page_number
                    st.rerun()

    st.markdown('</div>', unsafe_allow_html=True)

# =========================
# PREVIEW BOVEN
# =========================

st.markdown('<div class="section-card">', unsafe_allow_html=True)

if image_path is not None:
    st.image(str(image_path), use_container_width=True)
else:
    st.warning(f"Geen {file_type_label}-afbeelding gevonden.")

st.markdown(
    f'<div class="small-note">{file_type_label.capitalize()} {current_index + 1} blijft hierboven zichtbaar terwijl je hieronder leest.</div>',
    unsafe_allow_html=True
)

st.markdown('</div>', unsafe_allow_html=True)

st.write("")

# =========================
# AI PANEL ONDER
# =========================

st.markdown('<div class="section-card">', unsafe_allow_html=True)
st.markdown("### AI studiehulp")

explain_key = get_cache_key(hash_value, current_index, "explain", st.session_state.language)
simple_key = get_cache_key(hash_value, current_index, "simple", st.session_state.language)
keypoints_key = get_cache_key(hash_value, current_index, "keypoints", st.session_state.language)
questions_key = get_cache_key(hash_value, current_index, "questions", st.session_state.language)
chat_key = get_cache_key(hash_value, current_index, "chat", st.session_state.language)

if not current_text.strip():
    st.info(f"Deze {file_type_label} bevat geen leesbare tekst.")
elif image_path is None:
    st.info(f"Er is geen {file_type_label}-afbeelding beschikbaar voor Gemini.")
else:
    if explain_key not in st.session_state.ai_outputs:
        with st.spinner("Gemini bekijkt de pagina..."):
            st.session_state.ai_outputs[explain_key] = generate_short_explanation(
                current_text=current_text,
                deck_memory=deck_memory,
                page_number=current_index + 1,
                image_path=image_path
            )

    st.markdown("#### Korte uitleg")
    st.write(st.session_state.ai_outputs[explain_key])

    btn1, btn2 = st.columns(2)
    btn3, btn4 = st.columns(2)

    with btn1:
        if st.button("Probeer opnieuw", key=f"retry_{current_index}", use_container_width=True):
            with st.spinner("Gemini probeert opnieuw..."):
                st.session_state.ai_outputs[explain_key] = generate_short_explanation(
                    current_text=current_text,
                    deck_memory=deck_memory,
                    page_number=current_index + 1,
                    image_path=image_path
                )
            st.rerun()

    with btn2:
        if st.button("Leg simpeler uit", key=f"simple_{current_index}", use_container_width=True):
            with st.spinner("Gemini vereenvoudigt..."):
                st.session_state.ai_outputs[simple_key] = generate_simpler_explanation(
                    current_text=current_text,
                    deck_memory=deck_memory,
                    page_number=current_index + 1,
                    image_path=image_path
                )
            st.rerun()

    with btn3:
        if st.button("Wat onthouden?", key=f"remember_{current_index}", use_container_width=True):
            with st.spinner("Gemini kiest de kern..."):
                st.session_state.ai_outputs[keypoints_key] = generate_key_points(
                    current_text=current_text,
                    deck_memory=deck_memory,
                    page_number=current_index + 1,
                    image_path=image_path
                )
            st.rerun()

    with btn4:
        if st.button("3 oefenvragen", key=f"questions_{current_index}", use_container_width=True):
            with st.spinner("Gemini maakt vragen..."):
                st.session_state.ai_outputs[questions_key] = generate_questions(
                    current_text=current_text,
                    deck_memory=deck_memory,
                    page_number=current_index + 1,
                    image_path=image_path
                )
            st.rerun()

    if simple_key in st.session_state.ai_outputs:
        st.markdown("#### Simpeler uitgelegd")
        st.write(st.session_state.ai_outputs[simple_key])

    if keypoints_key in st.session_state.ai_outputs:
        st.markdown("#### Wat je moet onthouden")
        st.write(st.session_state.ai_outputs[keypoints_key])

    if questions_key in st.session_state.ai_outputs:
        st.markdown("#### Oefenvragen")
        st.write(st.session_state.ai_outputs[questions_key])

    st.markdown(f"#### Stel een vraag over deze {file_type_label}")

    with st.form(key=f"ask_form_{current_index}", clear_on_submit=True):
        user_question = st.text_input(
            "Typ je vraag",
            placeholder="Bijvoorbeeld: hoe werkt deze formule precies?"
        )
        ask_submit = st.form_submit_button("Verstuur vraag", use_container_width=True)

        if ask_submit and user_question.strip():
            with st.spinner("Gemini denkt na..."):
                chat_history = st.session_state.chat_answers.get(chat_key, [])
                answer = ask_about_page(
                    current_text=current_text,
                    deck_memory=deck_memory,
                    page_number=current_index + 1,
                    image_path=image_path,
                    question=user_question.strip(),
                    recent_chat_history=chat_history
                )
                chat_history.append({
                    "question": user_question.strip(),
                    "answer": answer
                })
                st.session_state.chat_answers[chat_key] = chat_history
            st.rerun()

    if chat_key in st.session_state.chat_answers:
        st.markdown("#### Vragen en antwoorden")
        for item in st.session_state.chat_answers[chat_key]:
            st.markdown(f"**Jij:** {item['question']}")
            st.write(item["answer"])
            st.markdown("---")

st.markdown('</div>', unsafe_allow_html=True)

# =========================
# NAVIGATIE ONDERAAN
# =========================

st.write("")
st.markdown("### Navigatie")

nav1, nav2, nav3 = st.columns([1, 1, 1])

with nav1:
    if st.button("⬅️ Vorige", use_container_width=True, disabled=(current_index == 0)):
        st.session_state.page_index -= 1
        st.rerun()

with nav2:
    st.markdown(
        f"<div style='text-align:center; padding-top:10px; font-weight:700;'>{current_index + 1} / {total_pages}</div>",
        unsafe_allow_html=True
    )

with nav3:
    if st.button("Volgende ➡️", use_container_width=True, disabled=(current_index >= total_pages - 1)):
        st.session_state.page_index += 1
        st.rerun()